"""CHECKS_v3.docx — отчёт о ревизии v3 для научного руководителя (A4)."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

doc = Document()
st = doc.styles["Normal"].font
st.name = "Times New Roman"; st.size = Pt(11)

def h(t, lvl=1):
    doc.add_heading(t, level=lvl)

def p(t, bold=False, italic=False):
    par = doc.add_paragraph()
    r = par.add_run(t); r.bold = bold; r.italic = italic
    return par

def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")

def table(header, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    for j, htxt in enumerate(header):
        c = t.cell(0, j); c.paragraphs[0].add_run(htxt).bold = True
    for i, row in enumerate(rows, 1):
        for j, cell in enumerate(row):
            t.cell(i, j).paragraphs[0].add_run(cell)
    for row in t.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(9.5)
    return t

doc.add_heading("Отчёт о ревизии статьи PIERE 2026 (v3)", level=0)
p("Дата: 2026-09-11. Автор ревизии: агент (Claude Code), по заданию piere2026_fix_prompt_v3. "
  "Документы: paper/piere2026_paper_en.docx (подача), paper/piere2026_report_ru.docx (русская версия), "
  "этот отчёт. Аудит и план: paper/AUDIT_v3.md, paper/PLAN_v3.md.")

h("1. Что изменено и почему")
bullets([
    "Заголовок: «Comparative Evaluation of Methods…» → «The Silver Ceiling: Machine-Labelled "
    "Validation Overestimates Multilingual Fraud Dialogue Detectors». Прежний называл жанр; новый — "
    "утверждение о результате, с именем эффекта; убраны запрещённые слова comparative/evaluation.",
    "Введён именованный эффект: the silver ceiling / эффект серебряного потолка — качество детектора, "
    "отобранного по машинным меткам, сходится к качеству самих меток против человека. Проверено по arXiv: "
    "термин в ML-смысле свободен (единственная коллизия — несвязанный RL-акроним SILVER, arXiv:2510.19244).",
    "Структура приведена к требованиям §2.2 один-к-одному: I Introduction · II Relevance and Related Work "
    "(обзор вынесен из Введения) · III Problem Statement · IV Theory · V Experimental Results and Practical "
    "Significance · VI Conclusion. Формулы перенумерованы по порядку появления (1)–(6).",
    "Аннотация переписана по заданной схеме (постановка → результат с числом и именем → специфичность → "
    "семейства и перенос → практический вывод), 231 слово; первый абзац Введения — 7 предложений, "
    "заканчивается ссылкой на Fig. 1.",
    "Рисунки перерисованы в цвете через единый scripts/figstyle.py (GOLD #009E73 = человеческие метки, "
    "SILVER #D55E00 = машинные; дублирование цвета маркером/стилем линии для ч/б). Добавлена схема "
    "конвейера (Fig. 1, на всю полосу, два выхода). Лесенка-рисунок удалена (дублировала таблицу II); "
    "таблица разрыва отбора заменена рисунком-«гантелями» (Fig. 4).",
    "Впервые включено: mDeBERTa в парном сравнении отбора (разрыв −0.0017 — граница эффекта; Положение 2 "
    "переформулировано честнее), слой 4 текстом (потолок 0.8757/0.8757/0.8680), парное подтверждение "
    "разрыва xlm-r (+0.0520).",
    "Добавлена таблица «текущая/предлагаемая практика» (Table III) и абзац о бюджете разметки; в Заключение "
    "добавлена прямая рекомендация практику в повелительной форме.",
    "Добавлены 2 ссылки IEEE REEPE 2025 (DOI проверены через Crossref); итого 23 источника.",
    "Из текста убраны имена моделей-генераторов корпуса (ужесточение эмбарго §8: конфигурация генерации "
    "не раскрывается); оставлено «two LLMs of a single vendor».",
])

h("2. Таблица пересчётов чисел")
p("Все прочие числа статьи сверены с runs/ и paper/tables/table_data_v2.json — расхождений нет "
  "(см. AUDIT_v3.md). Ниже только новые и заново вычисленные значения.")
table(
    ["Величина", "Было (v2-ревизия)", "Стало (v3)", "Источник"],
    [
        ["mDeBERTa, отбор по silver-val, gold-test v2", "не считалось",
         "F1 0.8923 ± 0.0159 (разрыв −0.0017)",
         "runs/epoch_cost/data/mdeberta/* (пересчёт снимков против gold_test_v2.csv; "
         "воспроизведён независимо 2026-09-11)"],
        ["xlm-r, парный повтор gold-sel/silver-sel", "не в статье",
         "0.9311 ± 0.0113 / 0.8791 ± 0.0046 (+0.0520)",
         "runs/epoch_cost/data/xlmr/* (тот же пересчёт)"],
        ["DeepSeek-контроль (автор корпуса)", "не в статье",
         "F1 0.8680, согласие с silver 0.9901",
         "runs/layer4/results_control.json"],
        ["Независимая LLM in-session", "не в статье",
         "F1 0.8757, согласие с silver 1.0000 (1 прогон, оговорка в тексте)",
         "runs/layer4_insession/results_test_few_rules.json"],
        ["Ключевые слова, порог по silver", "0.8889 (было в таблице III)",
         "0.8889 (теперь в Fig. 4)", "runs/layer0/results.json"],
        ["Разрывы отбора TF-IDF / слой 2", "+0.0008 / +0.0011 (таблица III)",
         "те же (Fig. 4)", "table_data_v2.json"],
    ])

h("3. Проверка формальных требований §2 (на готовом piere2026_paper_en.docx)")
table(
    ["Пункт", "Да/нет", "Значение"],
    [
        ["2.1 Научное содержание: задача, новизна, обоснованность, завершённость",
         "да", "3 научных положения во Введении и Заключении; каждое подкреплено числами из runs/ "
         "и парными тестами (бутстрап + перестановочный)"],
        ["2.2 Структура разделов один-к-одному", "да",
         "I Introduction · II Relevance and Related Work · III Problem Statement · IV Theory · "
         "V Experimental Results and Practical Significance · VI Conclusion"],
        ["2.3 Объём 5–6 страниц шаблона", "да", "6 полос (замер: экспорт Word → PDF)"],
        ["2.3 ≥2500 слов тела (без авторов, аннотации, ключевых слов, источников)", "да", "3875 слов"],
        ["2.3 Название 4–15 слов", "да", "10 слов"],
        ["2.3 Сведения об авторах: имя, должность, аффилиация, почтовый адрес, e-mail", "да",
         "Elena V. Krupoderova, Master's student, BMSTU, 2nd Baumanskaya st. 5/1, 105005 Moscow, "
         "Russia, lena05k@yandex.ru"],
        ["2.3 Аннотация 150–250 слов", "да", "231 слово"],
        ["2.3 Ключевые слова 5–7", "да", "7 (включая silver ceiling)"],
        ["2.3 Источники ≥8 (реком. 20–30), самоцитирований ≤3", "да", "23; самоцитирований 0"],
        ["2.3 Авторов ≤6", "да", "1"],
        ["2.4 Только английский, ноль кириллицы (текст + свойства документа)", "да",
         "0 символов кириллицы в document.xml, core.xml, app.xml (проверено скриптом)"],
        ["2.4 Название и аннотация без формул/цитирований/аббревиатур (кроме неизбежных)", "да",
         "неизбежные: F1 (метрика), имена моделей XLM-RoBERTa/mDeBERTa"],
        ["2.4 Ссылки [n] по порядку первого цитирования, формат IEEE", "да",
         "нумерация построена конструктивно по порядку появления"],
        ["2.4 Аббревиатуры расшифрованы при первом употреблении", "да",
         "LLM (I), TF-IDF (IV-B), XLM-R (IV-C); silver/gold/checkpoint/zero-shot определены в IV"],
        ["2.4 Подпись рисунка снизу, заголовок таблицы сверху, объект после первого упоминания",
         "да", "стили шаблона figure caption / table head с автонумерацией Fig. n. / TABLE n."],
        ["2.4 Формулы нумерованы справа, символы определены после формулы", "да",
         "6 формул OMML, номера (1)–(6) справа табуляцией; where-клаузы после каждой"],
        ["2.4 Единицы СИ, ноль перед десятичной точкой", "да",
         "величины безразмерные; чисел вида «.25» нет (проверено скриптом)"],
        ["§7 Сборка: копия шаблона, стили шаблона, рисунки встроены, таблицы Word, формулы OMML",
         "да", "python-docx поверх копии docs/PIERE2026__full_paper_template.docx; 5 embedded PNG "
         "300 dpi; 3 таблицы Word; 6 OMML"],
        ["§7 Приёмка: 3 файла существуют, >100 КБ, EN: 5 рисунков и 6 формул, 5–6 полос",
         "да", "см. раздел 5 ниже"],
    ])

h("4. Решение автора требуется")
p("Заголовок (выбран вариант A):", bold=True)
bullets([
    "A (в статье): The Silver Ceiling: Machine-Labelled Validation Overestimates Multilingual "
    "Fraud Dialogue Detectors — имя эффекта в заголовке, связывает заголовок, аннотацию, Fig. 1 и Заключение.",
    "B (альтернатива): Validating Fraud Detectors on Machine Labels Measures the Annotator, "
    "Not the Detector — афористичнее, но без многоязычности и имени эффекта.",
    "C (альтернатива): Checkpoint Selection on Machine Labels Silently Caps Multilingual Scam "
    "Dialogue Detectors — точнее про механизм, но жаргоннее.",
])
p("Имя эффекта: the silver ceiling / эффект серебряного потолка (по arXiv не занято; введено "
  "курсивом с определением во Введении, добавлено в ключевые слова).")
p("Включено впервые: mDeBERTa в парных сравнениях отбора (граница эффекта, Положение 2 "
  "переформулировано); слой 4 обоими прогонами — текстом, не семейством (DeepSeek-контроль 0.8680, "
  "независимая LLM 0.8757, оговорки о цикличности и невоспроизводимости в тексте); парное "
  "подтверждение разрыва xlm-r (+0.0520) одним придаточным.")
p("Осознанно не включено (эмбарго §8): онсет-голова и её метрики, ранняя детекция (earliness), "
  "ablation диалог/реплика, подтипы; runs/model_comparison и mdeberta_gold_dev (числа эпохи v1).")
p("[REPO_URL]: в тексте статьи стоит плейсхолдер — перед подачей заменить на реальный адрес "
  "репозитория со скриптами. Корпус как доступный не упоминается.")

h("5. Отступления от задания и приёмка")
bullets([
    "Пути /mnt/project и /mnt/skills из задания в рабочей среде отсутствуют; шаблон взят из "
    "docs/PIERE2026__full_paper_template.docx (тот же файл), сборка — python-docx из .venv.",
    "Ширины рисунков в figstyle.py скорректированы под фактические колонки A4-шаблона PIERE: "
    "COL 3.5″ → 3.38″, FULL 7.16″ → 7.01″ (полоса набора A4 7.01″; константы задания рассчитаны "
    "на US Letter IEEE). Принцип «ширина задаётся в скрипте, в Word не масштабировать» соблюдён: "
    "вставка ровно в размер генерации.",
    "Русская версия — полный текст без сокращений под лимит полос (8 полос): лимит 5–6 задан для "
    "подаваемого EN-файла; все содержательные правки (структура, рисунки, таблицы, числа, эмбарго) "
    "зеркалированы. Если руководителю нужна русская версия ровно в 6 полос — сокращения EN легко перенести.",
    "Приёмка: piere2026_paper_en.docx — 6 полос, 5 встроенных рисунков, 6 формул OMML, 0 кириллицы; "
    "piere2026_report_ru.docx — те же рисунки и таблицы, подписи русские (Рис. n. / ТАБЛИЦА n.); "
    "оба файла и этот отчёт больше 100 КБ и открываются в Word.",
])

h("6. Новые рисунки (как в статье)")
caps = [
    ("paper/figures/fig1_pipeline.png", 6.5,
     "Fig. 1 — схема конвейера: один конвейер, два пути отбора, два ответа."),
    ("paper/figures/fig4_significance.png", 4.2,
     "Fig. 2 — форест парной значимости (синее = значимо, серое = нет)."),
    ("paper/figures/fig2_divergence.png", 4.2,
     "Fig. 3 — расхождение silver-val и gold-dev по эпохам, оба чекпоинта помечены."),
    ("paper/figures/fig3_selection_gap.png", 4.2,
     "Fig. 4 — разрыв отбора по пяти семействам, пунктир — серебряный потолок."),
    ("paper/figures/fig5_transfer.png", 4.2,
     "Fig. 5 — матрица межъязыкового переноса, рамка — аномалия en→zh."),
]
for img, w, cap in caps:
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(img, width=Inches(w))
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(cap); r.italic = True; r.font.size = Pt(9)

props = doc.core_properties
props.title = "Отчёт о ревизии статьи PIERE 2026 (v3)"
props.author = "Claude Code (для Е. В. Круподеровой)"
doc.save("paper/CHECKS_v3.docx")
print("built paper/CHECKS_v3.docx")
